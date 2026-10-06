"""Career Quest Agent: a local four-stage interview game."""

import io
import importlib
import os
import uuid
import wave
from array import array
from pathlib import Path

import streamlit as st
from docx import Document
from openai import OpenAI
from pypdf import PdfReader

# Reload the preceding release once when an existing local server stays open.
for module_name, version in (("agents", 1), ("storage", 2), ("ui", 2)):
    module = importlib.import_module(module_name)
    if getattr(module, "QUEST_VERSION", 0) < version:
        importlib.reload(module)

from agents import (
    PASS_SCORE,
    QUESTIONS_PER_STAGE,
    STAGES,
    evaluate_agent,
    make_client,
    model_name,
    profile_agent,
    question_agent,
)
from storage import get_history, init_db, save_result
from gameplay import DIFFICULTIES, reward_summary
from speech import SPEECH_MODELS, SpeechInputError, transcribe_recording, vocabulary_hints
import server_settings
from server_settings import PROVIDERS, public_error, resolve_settings
from hosting import HostingError, data_dir, public_mode, speech_slot, visitor_id
from speech_model import SpeechServiceError, load_cpu_model, report_speech_failure
from ui import (
    STAGE_TITLES,
    navigate,
    render_brand,
    render_feedback,
    render_help,
    render_history,
    render_intro,
    render_profile,
    render_stage_route,
    render_training_plan,
    render_mission,
    render_rewards,
)

# Existing servers may retain the earlier settings module until restarted.
SettingsError = getattr(server_settings, "SettingsError", ValueError)


def extract_resume(upload):
    if upload is None:
        return ""
    suffix = Path(upload.name).suffix.lower()
    raw = upload.getvalue()
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("简历文件请控制在 5 MB 以内。")
    if suffix == ".txt":
        return raw.decode("utf-8-sig")[:10000]
    if suffix == ".docx":
        document = Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(parts)[:10000]
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages[:15])[:10000]
    raise ValueError("请上传 PDF、DOCX 或 TXT 格式简历。")


def speech_metrics(raw, transcript):
    fillers = sum(transcript.count(word) for word in ("嗯", "呃", "那个", "就是"))
    result = {"fillers": fillers, "duration": None, "silence_ratio": None}
    try:
        with wave.open(io.BytesIO(raw), "rb") as audio:
            rate = audio.getframerate()
            result["duration"] = round(audio.getnframes() / rate, 1)
            if audio.getsampwidth() == 2:
                silent = total = 0
                while True:
                    frame = audio.readframes(max(1, int(rate * 0.2)))
                    if not frame:
                        break
                    samples = array("h")
                    samples.frombytes(frame)
                    if not samples:
                        continue
                    rms = (sum(value * value for value in samples) / len(samples)) ** 0.5
                    silent += int(rms < 500)
                    total += 1
                result["silence_ratio"] = round(silent / total * 100) if total else None
    except (wave.Error, EOFError, ZeroDivisionError, ValueError):
        pass
    return result


@st.cache_resource(max_entries=1, show_spinner=False)
def load_whisper(model_size="small"):
    """Cache one multilingual CPU model, keyed by the chosen size."""
    model_dir = data_dir() / "models" if public_mode() else Path(__file__).with_name("models")
    return load_cpu_model(model_size, model_dir)


def start_stage(game, client, model, index):
    stage = STAGES[index]
    question = question_agent(client, model, game["job"], game["profile"], stage,
                              difficulty=game.get("difficulty", "标准挑战"))
    game.update(stage_index=index, question_index=0, question=question, evaluations=[],
                latest_feedback=None, stage_result=None,
                attempt=game.get("attempt", 0) + 1)


def text_client():
    return make_client(provider, api_key.strip(), base_url=settings.base_url)


def show_request_error(action, error):
    st.error(str(error) if isinstance(error, HostingError) else public_error(action, error))


st.set_page_config(page_title="Career Quest · 职业面试", page_icon=":material/work:",
                   layout="wide", initial_sidebar_state="expanded")
try:
    try:
        server_options = st.secrets.get("career_quest", {})
    except FileNotFoundError:
        if any(path.is_file() for path in (
            Path.cwd() / ".streamlit" / "secrets.toml",
            Path.home() / ".streamlit" / "secrets.toml",
        )):
            raise SettingsError("服务器配置文件无法读取，请管理员检查格式后重启。") from None
        server_options = {}
    settings = resolve_settings(server_options, os.environ)
except SettingsError as error:
    st.error(str(error))
    st.stop()
except (ValueError, TypeError):
    st.error("网站模型配置不完整或格式有误，请管理员检查服务器配置后重启。")
    st.stop()

is_public = public_mode()
if is_public and not settings.shared:
    st.error("网站正在配置面试服务，请稍后访问。管理员需在部署平台的 Secrets 中填写 career_quest 配置。")
    st.stop()
if is_public and settings.provider == "Ollama 本机免费":
    st.error("公开部署尚未配置可访问的文字模型，请管理员改用云端模型服务。")
    st.stop()
recognition_location = "服务器" if settings.shared or is_public else "本机"
init_db()
render_brand()
game = st.session_state.get("game")

default_id = st.session_state.setdefault(
    "default_player_id", uuid.uuid4().hex[:12] if settings.shared else "demo001"
)
if is_public:
    default_id = visitor_id(st.session_state)

with st.sidebar:
    st.caption("工作空间")
    navigation_labels = {
        "面试工作台": ":material/work:  面试工作台",
        "成长记录": ":material/trending_up:  成长记录",
        "使用说明": ":material/help_outline:  使用说明",
    }
    page = st.radio("工作空间导航", list(navigation_labels), key="workspace_page",
                    format_func=navigation_labels.get, label_visibility="collapsed")
    st.space("medium")
    st.caption("挑战进度")
    if game:
        st.write(game["job"])
        st.badge(game.get("difficulty", "标准挑战"), color="violet")
        stage_progress = (game["stage_index"] + (1 if game["stage_result"] else
                          game["question_index"] / QUESTIONS_PER_STAGE)) / len(STAGES)
        st.progress(stage_progress)
        st.caption(f"阶段 {game['stage_index'] + 1} / {len(STAGES)} · {STAGE_TITLES[game['stage_index']]}")
    else:
        st.markdown("**四道关卡，一条职业路线。**")
        st.caption("任务 · 解锁 · 经验 · 成就")
    st.space("medium")
    if settings.shared:
        provider, api_key, voice_key = settings.provider, settings.api_key, settings.voice_api_key
    else:
        with st.expander("模型与语音设置", icon=":material/tune:"):
            provider = st.selectbox("文字模型", PROVIDERS)
            api_key = st.text_input(
                "本机模式无需 API Key" if provider == "Ollama 本机免费" else f"{provider} API Key",
                type="password", disabled=provider == "Ollama 本机免费",
                help="只保留在当前会话中，不写入数据库。",
            )
            voice_key = st.text_input(
                "OpenAI 朗读 API Key（可选）", type="password",
                help="只用于面试官朗读；录音转文字无需此密钥。",
            )
            st.caption(f"当前模型：{model_name(provider)}")
            if provider == "阿里云百炼（免费额度）":
                st.caption("使用北京地域的密钥。建议在百炼控制台开启“免费额度用完即停”。")
            elif provider == "Ollama 本机免费":
                st.caption("请先在本机运行：ollama run qwen3:4b")
    with st.expander("语音识别设置", icon=":material/mic:"):
        available_modes = list(SPEECH_MODELS)
        default_speech_model = os.environ.get("CAREER_QUEST_WHISPER_MODEL", "base") if is_public else "small"
        speech_mode = st.selectbox("识别模式", available_modes, key="speech_mode",
                                  index=list(SPEECH_MODELS.values()).index(default_speech_model))
        custom_words = st.text_area("专业词汇（可选）", key="speech_words", persist_state="session",
                                    height=90, max_chars=120, placeholder="只填本段确实会说的术语，例如 Python、FastAPI；也可留空",
                                    help="最多采用 8 个短词。不自动加入岗位或简历技能；词汇提示也可能引入误识别。")
        use_vad = st.checkbox("过滤静音片段", value=True, key="speech_vad")
        st.caption(f"识别在{recognition_location}进行，不需要语音 API Key。small 首次下载约 486 MB，base 约 148 MB；small 更耗时。")
        st.caption("轻声内容被漏掉时，可关闭静音过滤后重试。")
    st.caption("Career Quest · 职业面试训练")
    st.caption("练习、复盘，准备你的下一次机会。")

local_text_model = provider == "Ollama 本机免费"
current_model = settings.model or model_name(provider)

with st.container(horizontal_alignment="center"):
    with st.container(width=1180, horizontal_alignment="left", gap="medium"):
        if page == "使用说明":
            render_help(recognition_location, public_demo=is_public)
            st.stop()

        if page == "成长记录":
            st.caption("职业准备 / 成长记录")
            st.title("看见进步，找到下一步。")
            st.write("回顾每次练习的表现，把反馈转化为更有针对性的准备。")
            lookup_default = game["player_id"] if game else st.session_state.get("last_player_id", default_id)
            if is_public:
                history_id = default_id
                st.caption("这里只显示本次会话的训练记录。关闭或刷新会话前请导出，公开体验版暂不支持跨会话找回。")
            else:
                history_id = st.text_input("存档编号", value=lookup_default, max_chars=40,
                                           key="history_lookup", help="输入之前练习时使用的编号。")
            st.session_state.last_player_id = history_id.strip()
            render_history(history_id.strip())
            st.stop()

        if game is None:
            render_intro()
            preparation, overview = st.columns([1.8, 1], gap="large")
            with preparation:
                with st.container(border=True):
                    st.subheader("创建你的职业挑战")
                    st.caption("设置目标岗位，我们会结合你的背景准备面试问题。")
                    with st.form("setup", border=False):
                        job = st.text_input("目标岗位", value="Python 开发工程师", max_chars=100,
                                            key="setup_job", persist_state="session",
                                            placeholder="例如：产品设计师、产品经理、Python 开发工程师")
                        difficulty = st.radio("挑战难度", list(DIFFICULTIES), key="setup_difficulty",
                                              help="首次练习推荐新手引导。各难度通关目标均为 70 分，但评分标准不同。")
                        st.caption("新手：基础问题与思路卡；标准：岗位与项目；进阶：取舍与情境追问。")
                        st.markdown("**简历信息** · 可选")
                        st.caption("提供项目与工作经历，让面试问题更贴合你。")
                        upload_tab, text_tab = st.tabs(["上传文件", "粘贴文字"])
                        with upload_tab:
                            resume_upload = st.file_uploader("上传简历", type=["pdf", "docx", "txt"],
                                                            help="支持 PDF、DOCX、TXT，建议不超过 5 MB。")
                        with text_tab:
                            resume_text = st.text_area("简历或个人经历", height=140,
                                                       key="setup_resume_text", persist_state="session",
                                                       placeholder="简要介绍你的技能、项目经历与工作背景。")
                        with st.expander("训练存档", icon=":material/bookmark:"):
                            if is_public:
                                player_id = default_id
                                st.caption("已为本次会话创建独立存档。结束练习后可在成长记录中导出。关闭或刷新页面可能结束当前会话。")
                            else:
                                player_id = st.text_input("存档编号",
                                                      value=st.session_state.get("last_player_id", default_id),
                                                      max_chars=40, key="setup_player_id", persist_state="session",
                                                      help="保留此编号，下次可查看之前的成长记录。")
                        st.caption("请使用脱敏简历。简历和回答文字将用于生成问题与练习反馈。")
                        start = st.form_submit_button("开始挑战", type="primary", width="stretch",
                                                       icon=":material/arrow_forward:")
                if start:
                    if (not local_text_model and not api_key.strip()) or not player_id.strip() or not job.strip():
                        st.error("请填写存档编号和目标岗位。" if settings.shared else
                                 "请填写目标岗位、存档编号，并在左侧设置中配置文字模型密钥。")
                    else:
                        try:
                            extracted = resume_text.strip() or extract_resume(resume_upload)
                            if resume_upload and not extracted:
                                st.warning("这份简历未读取到文字。可以改用粘贴文字的方式填写。")
                            client = text_client()
                            with st.spinner("正在分析你的背景，准备第一道面试问题…"):
                                profile = profile_agent(client, current_model, job.strip(), extracted)
                                candidate = {"player_id": player_id.strip(), "job": job.strip(),
                                             "profile": profile, "provider": provider,
                                             "difficulty": difficulty, "run_id": uuid.uuid4().hex}
                                start_stage(candidate, client, current_model, 0)
                            st.session_state.game = candidate
                            st.session_state.last_player_id = player_id.strip()
                            st.rerun()
                        except Exception as error:
                            show_request_error("启动失败", error)
                st.session_state.last_player_id = player_id.strip()
            with overview:
                render_training_plan()
                render_rewards(player_id.strip())
            render_history(player_id.strip(), compact=True)

        else:
            st.caption("CAREER QUEST / 挑战进行中")
            st.title(game["job"])
            st.write("完成本关任务，解锁下一位面试官。允许思考、修改和重试。")
            game.setdefault("run_id", uuid.uuid4().hex)
            reward_event = st.session_state.pop("quest_reward_event", None)
            if reward_event:
                st.toast(reward_event, icon=":material/stars:")
            if provider != game["provider"]:
                st.warning("本轮使用开始时选择的文字模型。请在设置中切回原模型，或结束本轮重新开始。")
            render_stage_route(game)
            stage = STAGES[game["stage_index"]]
            completed_count = game["question_index"] + (1 if game["stage_result"] else 0)
            st.progress((game["stage_index"] + completed_count / QUESTIONS_PER_STAGE) / len(STAGES))
            interview, feedback_panel = st.columns([1.8, 1], gap="large")

            with interview:
                if game["stage_result"]:
                    result = game["stage_result"]
                    with st.container(border=True):
                        st.subheader(f"{STAGE_TITLES[game['stage_index']]} · 关卡结算")
                        st.metric("本阶段平均分", result["score"])
                        st.badge("阶段已通关" if result["passed"] else "还有提升空间",
                                 color="green" if result["passed"] else "orange")
                        st.badge(f"本次获得 +{result.get('xp_earned', 0)} XP", color="violet", icon=":material/stars:")
                        if not result["passed"]:
                            st.info("本次完成经验已计入。先练习右侧建议，再重试本关；没有失败次数限制。")
                        st.markdown("**下一步练习**")
                        st.write(result["feedback"]["training_task"])
                        if result["passed"] and game["stage_index"] == len(STAGES) - 1:
                            st.success("本轮四个阶段已全部完成。前往成长记录，查看你的个人复盘。")
                            st.button("查看成长记录", type="primary", icon=":material/trending_up:",
                                      on_click=navigate, args=("成长记录",))
                        elif result["passed"]:
                            if st.button("进入下一关", type="primary", icon=":material/arrow_forward:",
                                         disabled=provider != game["provider"]):
                                try:
                                    with st.spinner("下一位面试官正在准备问题…"):
                                        start_stage(game, text_client(), current_model, game["stage_index"] + 1)
                                    st.rerun()
                                except Exception as error:
                                    show_request_error("生成下一关失败", error)
                        else:
                            if st.button("重试本关", type="primary", icon=":material/replay:",
                                         disabled=provider != game["provider"]):
                                try:
                                    with st.spinner("正在生成新的练习题…"):
                                        start_stage(game, text_client(), current_model, game["stage_index"])
                                    st.rerun()
                                except Exception as error:
                                    show_request_error("重试失败", error)

                else:
                    question_key = f"{game['run_id']}_{game['attempt']}_{game['stage_index']}_{game['question_index']}"
                    legacy_key = f"{game['attempt']}_{game['stage_index']}_{game['question_index']}"
                    for field in ("answer", "metrics", "speech"):
                        old_key, new_key = f"{field}_{legacy_key}", f"{field}_{question_key}"
                        if new_key not in st.session_state and old_key in st.session_state:
                            st.session_state[new_key] = st.session_state[old_key]
                    selected_voice_key = voice_key.strip() or (api_key.strip() if provider == "OpenAI" else "")
                    with st.container(border=True):
                        with st.container(horizontal=True, horizontal_alignment="distribute", vertical_alignment="center"):
                            st.subheader("当前任务 · 面试官提问")
                            st.badge(f"第 {game['question_index'] + 1} / {QUESTIONS_PER_STAGE} 题", color="gray")
                        with st.container(horizontal=True, vertical_alignment="center"):
                            st.image(str(Path(__file__).with_name("assets") / f"interviewer-{game['stage_index'] + 1}.svg"), width=64)
                            st.markdown(f"**{stage['name']} · 你的本关面试官**  \n{stage['persona']}")
                        st.caption(f"{STAGE_TITLES[game['stage_index']]} · {stage['focus']}")
                        st.write(game["question"])
                        if st.button("播放面试官语音", key=f"tts_{question_key}", icon=":material/volume_up:",
                                     disabled=settings.shared and not selected_voice_key):
                            if not selected_voice_key:
                                st.info("朗读尚未配置。可以直接阅读问题，继续录音回答。")
                            else:
                                try:
                                    with st.spinner("正在准备面试官语音…"):
                                        from hosting import api_request
                                        with api_request():
                                            speech = OpenAI(api_key=selected_voice_key, max_retries=0).audio.speech.create(
                                                model="gpt-4o-mini-tts", voice="alloy", input=game["question"]
                                            )
                                    st.session_state[f"speech_{question_key}"] = speech.content
                                except Exception as error:
                                    show_request_error("朗读失败", error)
                        if st.session_state.get(f"speech_{question_key}"):
                            st.caption("AI 合成语音")
                            st.audio(st.session_state[f"speech_{question_key}"], format="audio/mp3")
                        elif settings.shared and not selected_voice_key:
                            st.caption("本轮可阅读问题后直接录音作答。")

                    render_mission(game)
                    with st.container(border=True):
                        st.subheader("你的回答")
                        st.caption("录音 → 识别文字 → 确认提交，也可以直接输入回答。")
                        audio = st.audio_input("录制回答", key=f"audio_{question_key}",
                                               help="点击麦克风开始录音，再次点击结束。首次使用请允许麦克风权限。")
                        draft_exists = bool(st.session_state.get(f"answer_{question_key}", "").strip())
                        replace_draft = False
                        if audio and draft_exists:
                            replace_draft = st.checkbox("允许重新识别替换当前回答文字", key=f"replace_{question_key}")
                        if audio and st.button("识别我的录音", key=f"transcribe_{question_key}",
                                               icon=":material/graphic_eq:", disabled=draft_exists and not replace_draft):
                            try:
                                with st.spinner("正在识别录音；首次使用需要下载并加载模型…"):
                                    words = vocabulary_hints(custom_words)
                                    with speech_slot():
                                        transcription = transcribe_recording(audio.getvalue(), load_whisper(SPEECH_MODELS[speech_mode]),
                                                                            words, use_vad=use_vad)
                                st.session_state[f"answer_{question_key}"] = transcription["text"]
                                st.session_state[f"transcript_{question_key}"] = transcription
                                st.session_state[f"reviewed_{question_key}"] = False
                                st.session_state[f"metrics_{question_key}"] = speech_metrics(audio.getvalue(), transcription["text"])
                            except SpeechInputError as error:
                                st.warning(str(error))
                            except HostingError as error:
                                st.info(str(error))
                            except SpeechServiceError as error:
                                st.error(str(error))
                            except Exception as error:
                                report_speech_failure("audio-transcription", error)
                                st.error("录音未能完成识别，请管理员查看识别日志。可暂时输入文字继续练习。（诊断编号：S05）")
                        transcription = st.session_state.get(f"transcript_{question_key}")
                        if transcription:
                            for notice in transcription.get("warnings", []):
                                st.warning(notice)
                            st.caption("请核对专业术语、数字、人名和‘不 / 没有’等否定词。可回听上方录音，再修改下方文字。")
                            with st.expander("查看原始转写与待核对片段", icon=":material/hearing:"):
                                if transcription.get("retried"):
                                    st.caption("首次识别结果（供回听对照）")
                                    st.write(transcription["original_text"])
                                    st.caption("自动重识别后的当前结果")
                                st.write(transcription["text"])
                                for piece in transcription["review"]:
                                    st.write(f"{piece['start']}–{piece['end']} 秒：{piece['text']}")
                                st.caption("片段提示只是模型不确定性信号，不等于识别错误；没有提示也仍需核对。")
                        answer = st.text_area("确认回答内容", height=180, max_chars=5000, key=f"answer_{question_key}",
                                               persist_state="session",
                                               placeholder="识别后的文字会显示在这里。你也可以直接输入，并在提交前修改。")
                        reviewed = st.checkbox("我已核对转写，按当前文字提交", key=f"reviewed_{question_key}") if transcription else True
                        metrics = st.session_state.get(f"metrics_{question_key}")
                        if metrics:
                            with st.expander("查看表达统计", icon=":material/analytics:"):
                                st.caption(f"录音 {metrics['duration']} 秒 · 填充词约 {metrics['fillers']} 次 · "
                                           f"低音量片段约 {metrics['silence_ratio']}%")
                                st.caption("粗略统计只作练习参考，不作为表达能力的正式测评。")
                        if st.button("提交回答并评分", type="primary", width="stretch", icon=":material/arrow_forward:"):
                            if not answer.strip():
                                st.warning("请先录音识别，或直接输入回答。")
                            elif not reviewed:
                                st.warning("请先核对并确认转写文字，避免识别误差影响本题反馈。")
                            elif (not local_text_model and not api_key.strip()) or provider != game["provider"]:
                                st.warning("请在左侧设置中配置本轮使用的文字模型。")
                            else:
                                try:
                                    client = text_client()
                                    with st.spinner("正在分析你的回答，生成反馈…"):
                                        feedback = evaluate_agent(client, current_model, game["job"], game["profile"],
                                                                  stage, game["question"], answer.strip(),
                                                                  difficulty=game.get("difficulty", "标准挑战"))
                                        if game["question_index"] + 1 < QUESTIONS_PER_STAGE:
                                            next_question = question_agent(client, current_model, game["job"], game["profile"],
                                                                           stage, answer.strip(), feedback["gap"],
                                                                           difficulty=game.get("difficulty", "标准挑战"))
                                    evaluations = game["evaluations"] + [feedback]
                                    if game["question_index"] + 1 < QUESTIONS_PER_STAGE:
                                        game["evaluations"] = evaluations
                                        game["latest_feedback"] = feedback
                                        game["question_index"] += 1
                                        game["question"] = next_question
                                    else:
                                        score = round(sum(item["score"] for item in evaluations) / len(evaluations))
                                        result = {"score": score, "passed": score >= PASS_SCORE, "feedback": feedback}
                                        result_key = f"{game['run_id']}_{game['attempt']}_{game['stage_index']}"
                                        saved_feedback = dict(feedback, _quest={"difficulty": game.get("difficulty", "标准挑战"),
                                                                            "result_key": result_key})
                                        before = reward_summary(get_history(game["player_id"]))
                                        save_result(game["player_id"], game["job"], stage["name"], score,
                                                    result["passed"], saved_feedback, result_key=result_key)
                                        after = reward_summary(get_history(game["player_id"]))
                                        result["xp_earned"] = after["xp"] - before["xp"]
                                        new_badges = [badge for badge in after["badges"] if badge not in before["badges"]]
                                        st.session_state.quest_reward_event = f"完成关卡 +{result['xp_earned']} XP" + (
                                            " · 解锁 " + "、".join(new_badges) if new_badges else "")
                                        game["evaluations"] = evaluations
                                        game["latest_feedback"] = feedback
                                        game["stage_result"] = result
                                    st.rerun()
                                except Exception as error:
                                    show_request_error("评分失败", error)
            with feedback_panel:
                render_rewards(game["player_id"])
                render_feedback(game["latest_feedback"])
                render_profile(game["profile"])
                st.caption("面试评分用于训练参考。关注具体建议，持续优化自己的表达。")

            if st.button("结束本轮，返回首页", icon=":material/arrow_back:"):
                st.session_state.last_player_id = game["player_id"]
                del st.session_state.game
                st.rerun()

        st.caption("Career Quest · 每一次练习，都为下一次机会做准备。")
